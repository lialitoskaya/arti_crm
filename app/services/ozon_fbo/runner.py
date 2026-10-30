"""Выполнение прогона поставок FBO в рабочем потоке с состоянием в SQLite.

Прогон (расчёт, план, заявки, грузоместа) занимает минуты и требует ответов пользователя
(подтверждение плана, «всё равно сделать», подтверждение по артикулу). Поток пишет журнал,
строки результатов, вопрос и heartbeat в таблицы ozon_fbo_*, а ответы и флаг остановки
читает оттуда же, поэтому интерфейс может работать из любого процесса приложения.
"""
from __future__ import annotations

import logging
import os
import threading
import time
import traceback
from datetime import datetime
from pathlib import Path
from typing import Any, Callable, Optional

from app import ozon_fbo_repository as fbo_repo
from app.integrations.ozon_fbo.client import ApiError, AuthError, NetworkError, OzonFboApi, RunLog, get_credentials
from app.runtime_config import PROJECT_ROOT

from .cargo import run_cargo_orders_pipeline
from .config import FboSettings, parse_settings, use_settings
from .drafts import ArticleResult, RunOptions, RunOutcome
from .pipeline import article_summary, outcome_rows, print_summary, run_live_batch, run_pipeline
from .sheet import load_sheet, parse_manual_rows, write_results
from .store import FboStateStore
from .util import min_slot_date, norm, now_msk, safe_name, short

logger = logging.getLogger("arti_crm.ozon_fbo")

MODE_LABELS = {
    "trial": "Пробный прогон — заявки НЕ создаются, только расчёт «Факт» и слотов",
    "live": "Боевой — создать заявки и сразу передать грузоместа",
    "live_nocargo": "Боевой — создать заявки без грузомест",
    "cargo": "Дозаполнить грузоместа по уже созданным заявкам",
}
ANSWER_TIMEOUT_SECONDS = 2 * 60 * 60     # ждём ответа пользователя не дольше двух часов
ANSWER_POLL_SECONDS = 1.0
HEARTBEAT_SECONDS = 10.0
STALE_HEARTBEAT_SECONDS = 180            # без heartbeat дольше трёх минут прогон считается прерванным


class RunAbandonedError(RuntimeError):
    """Прогон уже завершён извне (помечен прерванным или снят), рабочему потоку продолжать нельзя."""


def log_base_dir() -> Path:
    raw = os.environ.get("OZON_FBO_LOG_DIR", "").strip()
    path = Path(raw) if raw else PROJECT_ROOT / "logs" / "ozon_fbo"
    if not path.is_absolute():
        path = PROJECT_ROOT / path
    return path


def load_active_settings() -> tuple[FboSettings, list[str]]:
    return parse_settings(fbo_repo.get_settings_raw())


def recover_interrupted_runs() -> list[int]:
    """Активные прогоны без heartbeat помечаются прерванными (гибель процесса, перезапуск хостинга)."""
    return fbo_repo.mark_stale_runs_interrupted(STALE_HEARTBEAT_SECONDS)


def _tone(res: ArticleResult, line, note: str) -> str:
    if line.fact == 0:
        return "zero"
    if res.manual:
        return "manual"
    return "ok" if line.fact else "none"


def rows_for_results(results: list[ArticleResult], mode: str) -> list[dict[str, Any]]:
    """Строки таблицы результатов (артикул, кластер, план, факт, примечание, оттенок) по результатам."""
    out: list[dict[str, Any]] = []
    for res in results:
        notes = outcome_rows(res, mode)[1]
        for ln in res.lines:
            note = ln.note or notes.get(ln.row, "") or ""
            out.append({"article": res.article, "cluster": ln.cluster, "plan": ln.plan, "fact": ln.fact,
                        "note": note, "tone": _tone(res, ln, note)})
        if not res.lines:
            out.append({"article": res.article, "cluster": "", "plan": 0, "fact": None,
                        "note": article_summary(res, mode) or res.note, "tone": "manual" if res.manual else "none"})
    return out


def summarize(outcome: RunOutcome, mode: str) -> dict[str, Any]:
    results = outcome.results
    return {
        "mode": mode,
        "articles": len(results),
        "created": sum(1 for r in results if r.created),
        "cargoes": sum(len(c.get("cargo_ids") or []) for r in results for c in r.cargoes.values()),
        "manual": sum(1 for r in results if r.manual),
        "skipped": sum(1 for r in results if r.skipped),
        "stopped": bool(outcome.stopped),
        "interrupted": outcome.interrupted is not None,
    }


class RunContext:
    """Связь рабочего потока с таблицей прогона: журнал, вопросы, остановка, heartbeat."""

    def __init__(self, run_id: int, *, sleep: Callable[[float], None] = time.sleep,
                 answer_timeout: float = ANSWER_TIMEOUT_SECONDS) -> None:
        self.run_id = run_id
        self.sleep = sleep
        self.answer_timeout = answer_timeout
        self.results: list[ArticleResult] = []
        self.mode = "trial"
        self.stop_seen = False      # пользователь нажал «Остановить» во время вопроса или между шагами

    def say(self, text: str = "") -> None:
        fbo_repo.append_run_log(self.run_id, text)

    def _ensure_active(self) -> dict[str, Any]:
        current = fbo_repo.get_run(self.run_id)
        if current is None or not current["is_active"]:
            raise RunAbandonedError(f"Прогон {self.run_id} уже завершён извне")
        return current

    def should_stop(self) -> bool:
        current = fbo_repo.get_run(self.run_id)
        if current is None or not current["is_active"]:
            return True
        if current["stop_requested"]:
            self.stop_seen = True
            return True
        return False

    def _take_answer(self) -> Optional[bool]:
        answer = fbo_repo.take_run_answer(self.run_id)
        if answer is None:
            return None
        self.say("  ответ пользователя: " + ("да" if answer == "yes" else "нет"))
        return answer == "yes"

    def wait_answer(self, kind: str, text: str) -> bool:
        fbo_repo.set_run_question(self.run_id, kind, text)
        deadline = time.monotonic() + self.answer_timeout
        while True:
            current = self._ensure_active()
            answered = self._take_answer()
            if answered is not None:
                return answered
            if current["stop_requested"]:
                self.stop_seen = True
            if current["stop_requested"] or time.monotonic() >= deadline:
                # вопрос снимается только если ответ не успел прийти; иначе честно учитываем ответ
                if not fbo_repo.clear_run_question(self.run_id):
                    answered = self._take_answer()
                    if answered is not None:
                        return answered
                    self._ensure_active()   # вопрос не снят и ответа нет — значит прогон закрыт извне
                self.say("  остановлено пользователем во время ожидания ответа — считаю ответ «нет»"
                         if current["stop_requested"] else "  ответ не получен — считаю ответ «нет»")
                return False
            self.sleep(ANSWER_POLL_SECONDS)

    def confirm(self, prompt: str, auto_confirm: bool, kind: str = "confirm") -> bool:
        if auto_confirm:
            self.say("  (подтверждено автоматически)")
            return True
        return self.wait_answer(kind, prompt)

    def ask_duplicates(self, prompt: str) -> bool:
        """Галочка «без подтверждения» этот вопрос не отвечает — разрешение на дубли всегда явное."""
        return self.wait_answer("duplicates", prompt)

    def on_result(self, res: ArticleResult, stage: str) -> None:
        if not any(existing is res for existing in self.results):
            self.results.append(res)
        fbo_repo.replace_run_rows(self.run_id, rows_for_results(self.results, self.mode))


def _heartbeat_loop(run_id: int, stop: threading.Event, interval: float = HEARTBEAT_SECONDS) -> None:
    while not stop.wait(interval):
        try:
            fbo_repo.heartbeat_run(run_id)
        except Exception:  # noqa: BLE001 — heartbeat не должен ронять прогон
            logger.warning("ozon_fbo heartbeat failed for run %s", run_id)


def start_run(run_id: int) -> threading.Thread:
    """Запускает прогон в фоновом потоке процесса. Состояние и результат — в БД."""
    thread = threading.Thread(target=execute_run, args=(run_id,), name=f"ozon-fbo-run-{run_id}", daemon=True)
    thread.start()
    return thread


def execute_run(run_id: int, *, api_factory: Optional[Callable[[RunLog], Any]] = None,
                store: Optional[FboStateStore] = None, sleep: Callable[[float], None] = time.sleep,
                now: Optional[datetime] = None, heartbeat: bool = True,
                answer_timeout: float = ANSWER_TIMEOUT_SECONDS, heartbeat_interval: float = HEARTBEAT_SECONDS) -> dict[str, Any]:
    """Выполняет прогон синхронно (вызывается из потока start_run или напрямую из тестов)."""
    run = fbo_repo.get_run(run_id)
    if run is None:
        raise ValueError(f"Прогон {run_id} не найден")
    ctx = RunContext(run_id, sleep=sleep, answer_timeout=answer_timeout)
    ctx.mode = run["mode"]
    settings, settings_errors = load_active_settings()
    log_dir = log_base_dir() / f"run_{run_id}"
    try:
        log_dir.mkdir(parents=True, exist_ok=True)
    except OSError:
        logger.warning("ozon_fbo: cannot create log dir for run %s", run_id)
    fbo_repo.mark_run_started(run_id, str(log_dir))
    stop_heartbeat = threading.Event()
    hb_thread: Optional[threading.Thread] = None
    if heartbeat:
        hb_thread = threading.Thread(target=_heartbeat_loop, args=(run_id, stop_heartbeat, heartbeat_interval),
                                     name=f"ozon-fbo-hb-{run_id}", daemon=True)
        hb_thread.start()
    status = "failed"
    error_text = ""
    abandoned = False
    summary: dict[str, Any] = {"mode": run["mode"]}
    result_name: Optional[str] = None
    result_bytes: Optional[bytes] = None
    api = None
    try:
        for err in settings_errors:
            ctx.say(f"Настройки: {err}")
        log = RunLog(log_dir, on_line=ctx.say)   # паузы лимитов и предупреждения клиента Ozon видны в журнале
        if api_factory is not None:
            api = api_factory(log)
        else:
            client_id, api_key = get_credentials()
            if not client_id or not api_key:
                raise RuntimeError("Ключи Ozon Seller API не заданы: нужны переменные среды OZON_CLIENT_ID и OZON_API_KEY.")
            api = OzonFboApi(client_id, api_key, log)
        with use_settings(settings):
            outcome, result_name, result_bytes = _execute_inner(run, ctx, api, store or fbo_repo.RepositoryStateStore(), log, now)
        summary = summarize(outcome, run["mode"])
        fbo_repo.replace_run_rows(run_id, rows_for_results(outcome.results, run["mode"]))
        if outcome.interrupted is not None and not isinstance(outcome.interrupted, KeyboardInterrupt):
            status, error_text = "failed", str(outcome.interrupted)
        elif outcome.stopped or ctx.stop_seen:
            status = "stopped"
        else:
            status = "finished"
    except RunAbandonedError as e:
        abandoned = True
        error_text = str(e)
        logger.warning("ozon_fbo run %s abandoned: %s", run_id, e)
    except AuthError as e:
        error_text = f"Ozon не принял ключи API или у ключа нет прав ({e.path}): {short(e.data, 300)}"
        ctx.say("ОШИБКА: " + error_text)
    except (ApiError, NetworkError, ValueError, FileNotFoundError, RuntimeError) as e:
        error_text = str(e)
        ctx.say("ОШИБКА: " + error_text)
    except Exception as e:  # noqa: BLE001 — любая другая ошибка фиксируется в прогоне, не роняя процесс
        error_text = f"Неожиданная ошибка: {type(e).__name__}: {e}"
        ctx.say("ОШИБКА: " + error_text)
        logger.error("ozon_fbo run %s crashed:\n%s", run_id, traceback.format_exc()[-3000:])
    finally:
        stop_heartbeat.set()
        if hb_thread is not None:
            hb_thread.join(timeout=heartbeat_interval + 1)
        if api is not None and hasattr(api, "close"):
            try:
                api.close()
            except Exception:  # noqa: BLE001
                pass
        current = fbo_repo.get_run(run_id)
        if abandoned or current is None or not current["is_active"]:
            # прогон уже закрыт извне (прерван или снят): итог другого процесса не перезаписываем
            status = str(current["status"]) if current else "interrupted"
        else:
            fbo_repo.finish_run(run_id, status, summary=summary, error_text=error_text,
                                result_file_name=result_name, result_file=result_bytes)
    return {"status": status, "summary": summary, "error_text": error_text}


def _execute_inner(run: dict[str, Any], ctx: RunContext, api: Any, store: FboStateStore, log: RunLog,
                   now: Optional[datetime]) -> tuple[RunOutcome, Optional[str], Optional[bytes]]:
    options_raw = run.get("options") or {}
    mode = run["mode"]
    auto_confirm = bool(options_raw.get("auto_confirm"))
    confirm_each = bool(options_raw.get("confirm_each"))
    force = bool(options_raw.get("force"))
    moment = now or now_msk()
    ctx.say(f"Режим: {MODE_LABELS.get(mode, mode)}")
    ctx.say(f"Запустил: {run.get('created_by_label') or '?'}")

    def confirm(prompt: str) -> bool:
        return ctx.confirm(prompt, auto_confirm, kind="confirm")

    def confirm_plan(prompt: str) -> bool:
        return ctx.confirm(prompt, auto_confirm, kind="plan")

    if mode == "cargo":
        keys = [str(k).strip() for k in options_raw.get("order_keys") or [] if str(k).strip()]
        ctx.say(f"Сейчас по Москве: {moment:%d.%m.%Y %H:%M}. Заявки: " + (", ".join("№" + k for k in keys) if keys else "все ожидающие"))
        outcome = run_cargo_orders_pipeline(api, store, keys or None, ctx.say, confirm, log, ctx.should_stop,
                                            confirm_each=confirm_each, auto_confirm=auto_confirm, on_result=ctx.on_result,
                                            confirm_plan=confirm_plan)
        ctx.say("Логи запросов Ozon доступны по ссылке «Логи запросов Ozon» в панели прогона.")
        return outcome, None, None

    source_kind = run["source_kind"]
    sheet = str(run.get("sheet_name") or "")
    wb = ws = layout = None
    if source_kind == "file":
        stored = fbo_repo.get_run_file(run["id"], "input")
        if stored is None:
            raise ValueError("Файл с планом не сохранён в прогоне")
        file_name, payload = stored
        wb, ws, layout, rows = load_sheet(payload, sheet)
        from openpyxl.utils import get_column_letter
        ctx.say(f"Файл: {file_name}, лист «{sheet}»")
        ctx.say(f"Заголовок в строке {layout.header_row}; план — колонка {get_column_letter(layout.col_plan)}, "
                f"факт — {get_column_letter(layout.col_fact)}, примечания — {get_column_letter(layout.col_note)}.")
    else:
        rows = parse_manual_rows(str(options_raw.get("manual_text") or ""))
        sheet = sheet or "ручной ввод"
        ctx.say(f"Ручной ввод: {len(rows)} строк, список «{sheet}»")
    ctx.say(f"Сейчас по Москве: {moment:%d.%m.%Y %H:%M}. Слот ищем не раньше {min_slot_date(moment):%d.%m.%Y}.")
    filt = {norm(a) for a in str(run.get("articles_filter") or "").split(",") if a.strip()}
    options = RunOptions(mode="live" if mode.startswith("live") else "trial", with_cargo=(mode == "live"), force=force,
                         articles=filt or None, auto_confirm=auto_confirm, confirm_each=confirm_each,
                         on_result=ctx.on_result, ask_duplicates=ctx.ask_duplicates, confirm_plan=confirm_plan)
    if options.mode == "live":
        outcome = run_live_batch(api, store, rows, sheet, options, moment, ctx.say, confirm, log, ctx.should_stop)
    else:
        outcome = run_pipeline(api, store, rows, sheet, options, moment, ctx.say, confirm, log, ctx.should_stop)
    if outcome.results:
        print_summary(outcome, options, ctx.say)
    result_name = result_bytes = None
    if source_kind == "file" and wb is not None and (outcome.facts or outcome.notes):
        stem = Path(str(run.get("input_file_name") or "план")).stem
        result_name = f"{safe_name(stem)}_результат.xlsx"
        result_bytes = write_results(wb, ws, layout, outcome.facts, outcome.notes)
        ctx.say(f"Результат: {result_name} (скачайте из раздела)")
    ctx.say("Логи запросов Ozon доступны по ссылке «Логи запросов Ozon» в панели прогона.")
    return outcome, result_name, result_bytes
