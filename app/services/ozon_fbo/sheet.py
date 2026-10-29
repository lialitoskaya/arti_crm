"""Таблица .xlsx и ручной ввод: чтение плана, запись результата, кластер «все», артикул или SKU."""
from __future__ import annotations

import io
import re
from dataclasses import dataclass
from typing import Any, BinaryIO, Optional

from .config import ALL_CLUSTERS_COUNTRY, ALL_CLUSTERS_WORDS, NOTE_HEADER, RED_FILL, cfg
from .util import norm


@dataclass
class Layout:
    header_row: int
    col_article: int
    col_cluster: int
    col_plan: int
    col_fact: int
    col_note: int


@dataclass
class TableRow:
    row: int
    article: str
    cluster: str
    plan: int
    hidden: bool


def _cell_text(ws, r, c) -> str:
    return norm(ws.cell(r, c).value)


def _column_empty(ws, col, from_row=1) -> bool:
    for r in range(from_row, ws.max_row + 1):
        v = ws.cell(r, col).value
        if v is not None and str(v).strip() != "":
            return False
    return True


def detect_layout(ws) -> Layout:
    """Находит строку заголовка и колонки: «Артикул», «Кластер», «Заявку на поставку» (план),
    следующая за планом — «Факт». Колонка для примечаний — первая пустая после «Факт»
    (в основной раскладке T/U/V это W: V — итог по артикулу, его не трогаем)."""
    header_row = None
    for r in range(1, min(ws.max_row, 15) + 1):
        for c in range(1, min(ws.max_column, 40) + 1):
            if _cell_text(ws, r, c) == "артикул":
                header_row = r
                break
        if header_row:
            break
    if not header_row:
        raise ValueError("Не нашёл строку заголовка с ячейкой «Артикул» в первых 15 строках листа.")
    cols = {}
    for c in range(1, ws.max_column + 1):
        cols[c] = _cell_text(ws, header_row, c)

    def find_col(pred):
        return next((c for c, t in cols.items() if pred(t)), None)

    col_article = find_col(lambda t: t == "артикул")
    col_cluster = find_col(lambda t: t == "кластер")
    # Точное совпадение важно: рядом есть колонка «В заявках на поставку» (уже поданные заявки),
    # её брать нельзя.
    col_plan = (find_col(lambda t: t == "заявку на поставку")
                or find_col(lambda t: t.startswith("заявку на поставку"))
                or find_col(lambda t: t.startswith("заявк") and "поставк" in t))
    if col_cluster is None:
        raise ValueError("В строке заголовка нет колонки «Кластер».")
    if col_plan is None:
        raise ValueError("В строке заголовка нет колонки «Заявку на поставку» (план).")
    col_fact = col_plan + 1
    col_note = col_fact + 2 if col_plan == 20 else col_fact + 1   # T/U/V → W
    while not _column_empty(ws, col_note):
        col_note += 1
    return Layout(header_row, col_article, col_cluster, col_plan, col_fact, col_note)


def read_rows(ws, layout: Layout) -> list[TableRow]:
    """Строки листа с планом > 0 (включая скрытые — они помечены hidden=True)."""
    rows = []
    last_article = ""
    for r in range(layout.header_row + 1, ws.max_row + 1):
        a = ws.cell(r, layout.col_article).value
        article = str(a).strip() if a is not None and str(a).strip() else last_article
        if a is not None and str(a).strip():
            last_article = article
        cluster = ws.cell(r, layout.col_cluster).value
        plan = ws.cell(r, layout.col_plan).value
        if isinstance(plan, str):
            try:
                plan = float(plan.replace(",", ".").strip())
            except ValueError:
                plan = None
        if not isinstance(plan, (int, float)) or plan <= 0:
            continue
        if not article or cluster is None or not str(cluster).strip():
            continue
        hidden = bool(ws.row_dimensions[r].hidden) if r in ws.row_dimensions else False
        rows.append(TableRow(r, article, str(cluster).strip(), int(round(plan)), hidden))
    return rows


def _as_stream(source: Any) -> Any:
    if isinstance(source, (bytes, bytearray)):
        return io.BytesIO(bytes(source))
    return source


def workbook_sheet_names(source: bytes | BinaryIO | str) -> list[str]:
    """Названия листов книги (для выбора листа перед прогоном)."""
    import openpyxl
    wb = openpyxl.load_workbook(_as_stream(source), read_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def load_sheet(source: bytes | BinaryIO | str, sheet_name: str):
    """Открывает книгу (значения вместо формул), возвращает (wb, ws, layout, rows)."""
    import openpyxl
    wb = openpyxl.load_workbook(_as_stream(source), data_only=True)
    if sheet_name not in wb.sheetnames:
        match = [s for s in wb.sheetnames if norm(s) == norm(sheet_name)]
        if not match:
            raise ValueError(f"Листа «{sheet_name}» нет в файле. Есть: {wb.sheetnames}")
        sheet_name = match[0]
    ws = wb[sheet_name]
    layout = detect_layout(ws)
    return wb, ws, layout, read_rows(ws, layout)


def write_results(wb, ws, layout: Layout, facts: dict, notes: dict) -> bytes:
    """Пишет «Факт» (красная заливка при нуле) и примечания в отдельную колонку.
    Возвращает содержимое копии книги; исходный файл не меняется."""
    from openpyxl.styles import PatternFill
    from openpyxl.utils import get_column_letter
    red = PatternFill(fill_type="solid", fgColor=RED_FILL, bgColor=RED_FILL)
    for r, fact in facts.items():
        cell = ws.cell(r, layout.col_fact)
        cell.value = int(fact)
        if int(fact) == 0:
            cell.fill = red
    if notes:
        header = ws.cell(layout.header_row, layout.col_note)
        if header.value in (None, ""):
            header.value = NOTE_HEADER
        letter = get_column_letter(layout.col_note)
        if not ws.column_dimensions[letter].width or ws.column_dimensions[letter].width < 40:
            ws.column_dimensions[letter].width = 60
        for r, text in notes.items():
            ws.cell(r, layout.col_note).value = text
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def resolve_products(api, keys: list[str]) -> dict[str, tuple[Optional[str], Optional[int]]]:
    """Введённое значение — артикул продавца или SKU Ozon (число) → (артикул как у Ozon, SKU).
    Не найденное → (None, None)."""
    offers = [k for k in keys if not str(k).strip().isdigit()]
    digits = [k for k in keys if str(k).strip().isdigit()]
    out: dict[str, tuple[Optional[str], Optional[int]]] = {}
    if offers:
        skus = api.skus_by_offer(offers)
        lower = {o.lower(): (o, s) for o, s in skus.items()}
        for k in offers:
            out[k] = (k, skus[k]) if k in skus else lower.get(k.lower(), (None, None))
    if digits:
        by_sku = api.offers_by_sku([int(k) for k in digits])
        for k in digits:
            oid = by_sku.get(int(k))
            out[k] = (oid, int(k)) if oid else (None, None)
    return out


def cluster_map_from(cluster_list: list[dict]) -> dict[str, tuple[int, str]]:
    result: dict[str, tuple[int, str]] = {}
    for c in cluster_list:
        result.setdefault(norm(c["name"]), (int(c["id"]), c["name"]))
    return result


def warehouse_cluster_map(cluster_list: list[dict]) -> dict[int, int]:
    """warehouse_id склада → id кластера (по fulfillments из /v2/cluster/list)."""
    out: dict[int, int] = {}
    for c in cluster_list:
        for wid in (c.get("warehouses") or {}):
            out.setdefault(int(wid), int(c["id"]))
    return out


def all_cluster_names(cluster_list: list[dict]) -> tuple[list[str], list[str]]:
    """Какие кластеры означает слово «все»: из настройки all_clusters или все кластеры России.
    Возвращает (названия как у Ozon, названия из настройки, которых у Ozon нет)."""
    by_norm = {norm(c["name"]): c["name"] for c in cluster_list}
    configured = cfg().all_clusters
    if configured:
        found = [by_norm[norm(n)] for n in configured if norm(n) in by_norm]
        missing = [n for n in configured if norm(n) not in by_norm]
        return found, missing
    return [c["name"] for c in cluster_list if norm(c.get("country")) == norm(ALL_CLUSTERS_COUNTRY)], []


def expand_all_clusters(rows: list, cluster_list: list[dict]) -> tuple[list, list[str]]:
    """Строки с кластером «все» разворачиваются в строку на каждый кластер из all_cluster_names с тем же
    количеством; кластеры, которые у артикула заданы отдельной строкой, не дублируются (явная строка
    важнее). Возвращает (строки, сообщения для отчёта)."""
    names, missing = all_cluster_names(cluster_list)
    messages = []
    if missing:
        messages.append(f"В настройке «все кластеры» есть названия, которых нет у Ozon: {missing}")
    explicit: dict[str, set] = {}
    for r in rows:
        if norm(r.cluster) not in ALL_CLUSTERS_WORDS:
            explicit.setdefault(norm(r.article), set()).add(norm(r.cluster))
    out = []
    for r in rows:
        if norm(r.cluster) not in ALL_CLUSTERS_WORDS:
            out.append(r)
            continue
        added = [n for n in names if n and norm(n) not in explicit.get(norm(r.article), set())]
        if not added:
            messages.append(f"стр.{r.row} {r.article}: «{r.cluster}» — список кластеров пуст, строка пропущена")
            continue
        messages.append(f"стр.{r.row} {r.article}: «{r.cluster}» → {len(added)} кластеров по {r.plan} шт.")
        for n in added:
            out.append(TableRow(r.row, r.article, n, r.plan, r.hidden))
    return out, messages


def parse_manual_rows(text: str) -> list[TableRow]:
    """Ручной ввод: по строке на кластер — «артикул; кластер; количество» (разделитель «;» или табуляция,
    как при копировании из таблицы). Строка-заголовок «Артикул …» пропускается."""
    rows = []
    for n, raw in enumerate(str(text or "").splitlines(), 1):
        line = raw.strip()
        if not line:
            continue
        parts = [p.strip() for p in re.split(r"\t|;", line)]
        if len(parts) < 3:
            raise ValueError(f"строка {n}: нужно три значения — артикул; кластер; количество (через «;» "
                             f"или табуляцию), а не {line!r}")
        article, cluster, qty = parts[0], parts[1], parts[2].replace(",", ".")
        if norm(article) == "артикул":
            continue
        if not article or not cluster:
            raise ValueError(f"строка {n}: пустой артикул или кластер: {line!r}")
        try:
            plan = int(round(float(qty)))
        except ValueError:
            raise ValueError(f"строка {n}: количество должно быть числом, а не {parts[2]!r}") from None
        if plan <= 0:
            continue
        rows.append(TableRow(n, article, cluster, plan, False))
    return rows
