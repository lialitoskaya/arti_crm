"""Поставки FBO: выполнение прогона с состоянием в SQLite — журнал, строки результатов, вопрос и ответ,
остановка, файл результата, ошибки ключей и forward recovery. Ozon подменён, сети нет."""
from __future__ import annotations

import io
import sys
import unittest
from pathlib import Path
from unittest import mock

_TESTS_DIR = str(Path(__file__).resolve().parent)
if _TESTS_DIR not in sys.path:
    sys.path.insert(0, _TESTS_DIR)

import test_regression_foundation as foundation  # noqa: E402
from test_ozon_fbo_core import DAYS, NOW, Sheet  # noqa: E402
from test_ozon_fbo_pipeline import BatchApi, OrdersApi, PipelineApi, api_with_today_orders  # noqa: E402

from app import db  # noqa: E402
from app import ozon_fbo_repository as fbo_repo  # noqa: E402
from app.integrations.ozon_fbo.client import AuthError  # noqa: E402
from app.services.ozon_fbo import runner  # noqa: E402


class _NoSleep:
    def __call__(self, seconds: float) -> None:
        return None


class OzonFboRunnerTests(unittest.TestCase):
    def setUp(self) -> None:
        foundation._remove_test_runtime_files()
        db.init_db()

    def _run(self, **over):
        params = dict(mode="trial", source_kind="manual", sheet_name="тест", articles_filter="",
                      options={"manual_text": "AT-P-P3; Тверь; 40\nAT-P-P3; Казань; 20"},
                      created_by_user_id=None, created_by_label="admin")
        params.update(over)
        return fbo_repo.create_run(**params)

    def _execute(self, run_id: int, api, sleep=None, **kw):
        with mock.patch("time.sleep", lambda s: None):
            return runner.execute_run(run_id, api_factory=lambda log: api, sleep=sleep or _NoSleep(), now=NOW,
                                      heartbeat=False, **kw)

    def test_trial_manual_run_finishes_with_log_rows_and_summary(self) -> None:
        run = self._run()
        api = PipelineApi({4041: 9}, DAYS)
        result = self._execute(run["id"], api)
        self.assertEqual(result["status"], "finished")
        stored = fbo_repo.get_run(run["id"])
        self.assertEqual((stored["status"], stored["is_active"], stored["summary"]["articles"], stored["summary"]["created"]),
                         ("finished", False, 1, 0))
        self.assertTrue(stored["log_dir"].endswith(f"run_{run['id']}"))
        text = fbo_repo.get_run_log_text(run["id"])
        self.assertIn("Режим: Пробный прогон", text)
        self.assertIn("ИТОГ ПРОБНОГО ПРОГОНА", text)
        self.assertIn("Запустил: admin", text)
        rows = fbo_repo.get_run_rows(run["id"])
        self.assertEqual([(r["cluster"], r["plan"], r["fact"], r["tone"]) for r in rows],
                         [("Тверь", 40, 40, "ok"), ("Казань", 20, 0, "zero")])
        self.assertIn("ЗАЯВКА НЕ СОЗДАНА", rows[0]["note"])
        self.assertEqual(api.supply_calls, [])
        self.assertEqual(fbo_repo.get_article_state("тест", "AT-P-P3")["facts"], {"Тверь": 40, "Казань": 0})
        self.assertTrue((Path(runner.log_base_dir()) / f"run_{run['id']}").exists(), "дампы запросов пишутся в каталог прогона")

    def test_live_run_waits_for_plan_answer_in_database(self) -> None:
        run = self._run(mode="live")
        api = BatchApi({}, DAYS)
        seen: list[dict] = []

        def sleep(seconds: float) -> None:
            current = fbo_repo.get_run(run["id"])
            seen.append(current)
            if current["question_kind"]:
                fbo_repo.answer_run(run["id"], "yes")

        result = self._execute(run["id"], api, sleep=sleep)
        self.assertEqual(result["status"], "finished")
        self.assertTrue(any(s["status"] == "awaiting_answer" and s["question_kind"] == "plan"
                            and s["question_text"].startswith("ПЛАН ПОСТАВОК") for s in seen), seen)
        stored = fbo_repo.get_run(run["id"])
        self.assertEqual((stored["status"], stored["question_kind"], stored["summary"]["created"], stored["summary"]["cargoes"]),
                         ("finished", None, 1, 2))
        self.assertEqual(len(api.supply_calls), 1)
        self.assertIn("ответ пользователя: да", fbo_repo.get_run_log_text(run["id"]))
        rows = fbo_repo.get_run_rows(run["id"])
        self.assertTrue(all(r["tone"] == "ok" and "заявка №" in r["note"] for r in rows), rows)

    def test_declined_plan_and_duplicates_question(self) -> None:
        run = self._run(mode="live", options={"manual_text": "AT-P-P3; Тверь; 40\nAT-P-P3; Казань; 20"})
        api = api_with_today_orders()
        kinds: list[str] = []

        def sleep(seconds: float) -> None:
            current = fbo_repo.get_run(run["id"])
            if current["question_kind"] and current["question_kind"] not in kinds:
                kinds.append(current["question_kind"])
                fbo_repo.answer_run(run["id"], "no")

        result = self._execute(run["id"], api, sleep=sleep)
        self.assertEqual(result["status"], "finished")
        self.assertEqual(kinds, ["duplicates", "plan"], "сначала вопрос про дубли, потом общий план")
        self.assertEqual(api.supply_calls, [])
        text = fbo_repo.get_run_log_text(run["id"])
        self.assertIn("решение пользователя: пропустить", text)
        self.assertIn("Создание заявок отменено пользователем", text)

    def test_auto_confirm_does_not_answer_duplicates(self) -> None:
        run = self._run(mode="live", options={"manual_text": "AT-P-P3; Тверь; 40\nAT-P-P3; Казань; 20", "auto_confirm": True})
        api = api_with_today_orders()
        kinds: list[str] = []

        def sleep(seconds: float) -> None:
            current = fbo_repo.get_run(run["id"])
            if current["question_kind"] and current["question_kind"] not in kinds:
                kinds.append(current["question_kind"])
                fbo_repo.answer_run(run["id"], "yes")

        self._execute(run["id"], api, sleep=sleep)
        self.assertEqual(kinds, ["duplicates"], "общий план подтверждён автоматически, дубли — только явно")
        self.assertEqual(sorted(api.supply_calls[0][1]), [4041, 4072])

    def test_stop_while_waiting_for_answer(self) -> None:
        run = self._run(mode="live")
        api = BatchApi({}, DAYS)

        def sleep(seconds: float) -> None:
            if fbo_repo.get_run(run["id"])["question_kind"]:
                fbo_repo.request_stop(run["id"])

        result = self._execute(run["id"], api, sleep=sleep)
        self.assertEqual(result["status"], "stopped")
        stored = fbo_repo.get_run(run["id"])
        self.assertEqual((stored["status"], stored["question_kind"], stored["is_active"]), ("stopped", None, False))
        self.assertEqual(api.supply_calls, [])
        self.assertIn("считаю ответ «нет»", fbo_repo.get_run_log_text(run["id"]))

    def test_answer_timeout_counts_as_no(self) -> None:
        run = self._run(mode="live")
        api = BatchApi({}, DAYS)
        clock = {"t": 0.0}

        def fake_monotonic() -> float:
            clock["t"] += 100.0
            return clock["t"]

        with mock.patch.object(runner.time, "monotonic", fake_monotonic):
            result = self._execute(run["id"], api, answer_timeout=50.0)
        self.assertEqual(result["status"], "finished")
        self.assertEqual(api.supply_calls, [])
        self.assertIn("ответ не получен", fbo_repo.get_run_log_text(run["id"]))

    def test_stop_between_articles_marks_stopped(self) -> None:
        run = self._run()
        fbo_repo.request_stop(run["id"])
        result = self._execute(run["id"], PipelineApi({}, DAYS))
        self.assertEqual(result["status"], "stopped")
        self.assertIn("Остановлено пользователем", fbo_repo.get_run_log_text(run["id"]))

    def test_file_source_produces_result_workbook(self) -> None:
        payload = Sheet.make_book()
        run = self._run(source_kind="file", sheet_name="Поставка 11.09", options={}, input_file_name="Поставки.xlsx", input_file=payload)
        api = PipelineApi({4041: 9}, DAYS)
        result = self._execute(run["id"], api)
        self.assertEqual(result["status"], "finished")
        stored = fbo_repo.get_run(run["id"])
        self.assertTrue(stored["has_result_file"])
        name, data = fbo_repo.get_run_file(run["id"], "result")
        self.assertEqual(name, "Поставки_результат.xlsx")
        import openpyxl
        ws = openpyxl.load_workbook(io.BytesIO(data))["Поставка 11.09"]
        self.assertEqual((ws["U4"].value, ws["U5"].value), (10, 0), "факт записан в копию таблицы")
        self.assertIn("Скрытые фильтром строки", fbo_repo.get_run_log_text(run["id"]))
        rows = fbo_repo.get_run_rows(run["id"])
        self.assertTrue(any(r["article"] == "AT-L-A4" and r["tone"] == "manual" for r in rows), "AT-L-A4 без SKU — ручная проверка")

    def test_cargo_mode_fills_selected_orders(self) -> None:
        run = self._run(mode="cargo", source_kind="orders", sheet_name="грузоместа",
                        options={"order_keys": ["2000070409946"]})
        api = OrdersApi()

        def sleep(seconds: float) -> None:
            if fbo_repo.get_run(run["id"])["question_kind"]:
                fbo_repo.answer_run(run["id"], "yes")

        result = self._execute(run["id"], api, sleep=sleep)
        self.assertEqual(result["status"], "finished")
        self.assertEqual([s for s, _ in api.cargo_calls], [7002, 7003])
        stored = fbo_repo.get_run(run["id"])
        self.assertEqual((stored["summary"]["cargoes"], stored["summary"]["articles"]), (6, 1))
        rows = fbo_repo.get_run_rows(run["id"])
        self.assertEqual([(r["cluster"], r["fact"]) for r in rows], [("НОГИНСК_РФЦ", 110), ("ПЕРМЬ_РФЦ", 12)])
        self.assertEqual(fbo_repo.get_order_state("2000070409946")["cargoes"]["7002"]["cargo_ids"], [7001, 7002, 7003])

    def test_auth_error_marks_run_failed_without_leaking_keys(self) -> None:
        run = self._run()

        class Broken(PipelineApi):
            def cluster_list(self):
                raise AuthError("/v2/cluster/list", 403, {"code": 7, "message": "Api-Key is missing a required role"})

        result = self._execute(run["id"], Broken({}, DAYS))
        self.assertEqual(result["status"], "failed")
        stored = fbo_repo.get_run(run["id"])
        self.assertIn("Ozon не принял ключи", stored["error_text"])
        self.assertFalse(stored["is_active"])
        self.assertIn("ОШИБКА", fbo_repo.get_run_log_text(run["id"]))

    def test_missing_keys_fail_fast(self) -> None:
        run = self._run()
        with mock.patch("time.sleep", lambda s: None):
            result = runner.execute_run(run["id"], sleep=_NoSleep(), now=NOW, heartbeat=False)
        self.assertEqual(result["status"], "failed")
        self.assertIn("Ключи Ozon", fbo_repo.get_run(run["id"])["error_text"])
        self.assertEqual([], foundation._NETWORK_ATTEMPTS)

    def test_unexpected_exception_is_recorded(self) -> None:
        run = self._run()

        class Crashing(PipelineApi):
            def cluster_list(self):
                raise KeyError("boom")

        result = self._execute(run["id"], Crashing({}, DAYS))
        self.assertEqual(result["status"], "failed")
        self.assertIn("Неожиданная ошибка", fbo_repo.get_run(run["id"])["error_text"])

    def test_recover_interrupted_runs(self) -> None:
        run = self._run()
        fbo_repo.mark_run_started(run["id"], "")
        self.assertEqual(runner.recover_interrupted_runs(), [])
        fbo_repo.update_run(run["id"], heartbeat_at="2020-01-01 00:00:00")
        self.assertEqual(runner.recover_interrupted_runs(), [run["id"]])
        self.assertEqual(fbo_repo.get_run(run["id"])["status"], "interrupted")

    def test_settings_from_database_apply_to_run(self) -> None:
        fbo_repo.save_settings_raw({"drop_off_name": "ТЕСТ_ТОЧКА", "default_cargo_capacity": 10})
        run = self._run(mode="live", options={"manual_text": "AT-P-P3; Тверь; 40", "auto_confirm": True})

        class DropApi(PipelineApi):
            def drop_off_point(self, name):
                self.drop_off_asked = name
                return {"warehouse_id": 1, "warehouse_type": "DELIVERY_POINT"}

        api = DropApi({}, DAYS)
        self._execute(run["id"], api)
        self.assertEqual(api.drop_off_asked, "ТЕСТ_ТОЧКА")
        self.assertEqual(len(api.cargo_calls[0][1]), 4, "коробы по 10 штук из настроек")

    def test_start_run_executes_in_background_thread(self) -> None:
        run = self._run()
        api = PipelineApi({}, DAYS)
        with mock.patch.object(runner, "execute_run") as execute:
            thread = runner.start_run(run["id"])
            thread.join(timeout=5)
        execute.assert_called_once_with(run["id"])
        self.assertFalse(thread.is_alive())
        del api


if __name__ == "__main__":
    unittest.main()
