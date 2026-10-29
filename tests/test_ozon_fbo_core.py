"""Поставки FBO: бизнес-правила, предохранители, расчёт артикула, боевой поток и клиент Ozon без сети.

Characterization-тесты перенесены из desktop-приложения «Поставки FBO Ozon» (test_ozon_fbo.py);
поведение правил сохранено один к одному."""
from __future__ import annotations

import io
import sys
import time
import unittest
from datetime import date, datetime, timedelta
from pathlib import Path
from unittest import mock

_TESTS_DIR = str(Path(__file__).resolve().parent)
if _TESTS_DIR not in sys.path:
    sys.path.insert(0, _TESTS_DIR)

import test_regression_foundation as foundation  # noqa: E402,F401

from app.integrations.ozon_fbo import client as api_client  # noqa: E402
from app.integrations.ozon_fbo.client import ApiError, OzonFboApi, RunLog, parse_order_ids, parse_orders, slot_days  # noqa: E402
from app.services.ozon_fbo import config as cfgmod  # noqa: E402
from app.services.ozon_fbo.cargo import build_cargoes, build_item_cargoes, cargo_qty, validate_order_cargoes  # noqa: E402
from app.services.ozon_fbo.config import MSK, FboSettings, parse_settings, settings_to_dict, use_settings  # noqa: E402
from app.services.ozon_fbo.drafts import ArticleResult, ClusterLine, process_article  # noqa: E402
from app.services.ozon_fbo.orders import order_warehouses  # noqa: E402
from app.services.ozon_fbo.pipeline import create_orders_batch, outcome_rows  # noqa: E402
from app.services.ozon_fbo.rules import (  # noqa: E402
    CheckResult, SupplyLine, candidates, cargo_layout, earliest_slot_date, find_max_accepted, layout_text,
    pick_latest_slot, validate_before_create,
)
from app.services.ozon_fbo.sheet import load_sheet, parse_manual_rows, write_results  # noqa: E402
from app.services.ozon_fbo.store import InMemoryStateStore, save_state  # noqa: E402
from app.services.ozon_fbo.util import min_slot_date, utc_to_msk_text  # noqa: E402


def capacity_checker(capacity, with_hint=False):
    """Кластер принимает не больше capacity штук. with_hint — Ozon сам урезает до capacity."""
    calls = []

    def check(q):
        calls.append(q)
        if q <= capacity:
            return CheckResult(True)
        return CheckResult(False, trimmed_to=capacity if with_hint else None, reason="нет слотов")
    return check, calls


class Candidates(unittest.TestCase):
    def test_lists(self):
        self.assertEqual(candidates(40), [40, 35, 30, 25, 20, 15, 10])
        self.assertEqual(candidates(38), [38, 35, 30, 25, 20, 15, 10])
        self.assertEqual(candidates(45), [45, 40, 35, 30, 25, 20, 15, 10])
        self.assertEqual(candidates(12), [12, 10])
        self.assertEqual(candidates(10), [10])
        self.assertEqual(candidates(8), [8])
        self.assertEqual(candidates(0), [])


class FindMaxAccepted(unittest.TestCase):
    cases = [  # (план, вместимость кластера, ожидаемый факт)
        (40, 50, 40), (40, 40, 40), (40, 38, 35), (40, 23, 20), (40, 14, 10), (40, 9, 0), (40, 0, 0),
        (45, 45, 45), (38, 37, 35), (10, 10, 10), (10, 9, 0), (12, 11, 10),
        (8, 8, 8), (8, 7, 0), (100, 63, 60), (20, 15, 15), (15, 12, 10), (70, 70, 70),
    ]

    def test_without_hint(self):
        for plan, cap, expected in self.cases:
            check, calls = capacity_checker(cap)
            got, history = find_max_accepted(plan, check)
            self.assertEqual(got, expected, f"план {plan}, вместимость {cap}")
            self.assertEqual(len(history), len(calls))
            self.assertEqual(calls[0], plan, "первая проверка — по плану")
            bound = 1 + (len(candidates(plan)) - 1).bit_length()
            self.assertLessEqual(len(calls), bound, f"слишком много проверок для плана {plan}: {calls}")
            if plan == 40:
                self.assertLessEqual(len(calls), 4)

    def test_with_hint(self):
        for plan, cap, expected in self.cases:
            check, calls = capacity_checker(cap, with_hint=True)
            got, _ = find_max_accepted(plan, check)
            self.assertEqual(got, expected, f"план {plan}, вместимость {cap}, проверки {calls}")
            self.assertLessEqual(len(calls), 2, f"с подсказкой Ozon нужно не больше 2 проверок: {calls}")

    def test_hint_below_ten_excludes(self):
        check, calls = capacity_checker(9, with_hint=True)
        got, _ = find_max_accepted(40, check)
        self.assertEqual(got, 0)
        self.assertEqual(calls, [40])

    def test_hint_verified_not_trusted_blindly(self):
        def check(q):
            if q > 23:
                return CheckResult(False, trimmed_to=23)
            return CheckResult(q <= 12)
        got, history = find_max_accepted(40, check)
        self.assertEqual(got, 10)
        self.assertEqual([q for q, _ in history][:2], [40, 20])

    def test_zero_plan(self):
        got, history = find_max_accepted(0, lambda q: CheckResult(True))
        self.assertEqual((got, history), (0, []))


class Cargo(unittest.TestCase):
    def test_layout(self):
        self.assertEqual(cargo_layout("AT-P-P3", 110), [50, 50, 10])
        self.assertEqual(cargo_layout("AT-P-P3", 50), [50])
        self.assertEqual(cargo_layout("AT-P-P3", 35), [35])
        self.assertEqual(cargo_layout("AT-P-P3", 0), [])
        self.assertEqual(cargo_layout("AT-L-AM-S", 12), [5, 5, 2])
        self.assertEqual(cargo_layout("at-l-am-m", 5), [5])
        self.assertEqual(cargo_layout("AT-L-AM-P", 20), [5, 5, 5, 5])


class Slots(unittest.TestCase):
    days = [
        {"date_in_timezone": "2026-09-15", "timeslots": [{"from_in_timezone": "2026-09-15T10:00:00", "to_in_timezone": "2026-09-15T11:00:00"}]},
        {"date_in_timezone": "2026-09-20", "timeslots": [
            {"from_in_timezone": "2026-09-20T19:00:00", "to_in_timezone": "2026-09-20T20:00:00"},
            {"from_in_timezone": "2026-09-20T09:00:00", "to_in_timezone": "2026-09-20T10:00:00"}]},
        {"date_in_timezone": "2026-09-25", "timeslots": []},
        {"date_in_timezone": "2026-09-18", "timeslots": [{"from_in_timezone": "2026-09-18T23:00:00", "to_in_timezone": "2026-09-19T00:00:00"}]},
    ]

    def test_latest_slot(self):
        slot = pick_latest_slot(self.days, date(2026, 9, 18))
        self.assertEqual(slot, {"date": "2026-09-20", "from": "2026-09-20T19:00:00", "to": "2026-09-20T20:00:00"})

    def test_not_before_filters(self):
        self.assertIsNone(pick_latest_slot(self.days, date(2026, 9, 21)))
        self.assertEqual(earliest_slot_date(self.days, date(2026, 9, 16)), "2026-09-18")
        self.assertIsNone(earliest_slot_date(self.days, date(2026, 9, 26)))

    def test_min_slot_date(self):
        now = datetime(2026, 9, 11, 23, 30, tzinfo=MSK)
        self.assertEqual(min_slot_date(now), date(2026, 9, 18))

    def test_slot_days_from_response(self):
        resp = {"result": {"drop_off_warehouse_timeslots": {"days": self.days}}}
        self.assertEqual(len(slot_days(resp)), 4)
        self.assertEqual(slot_days({}), [])


class Safeguards(unittest.TestCase):
    now = datetime(2026, 9, 11, 12, 0, tzinfo=MSK)
    ok_slot = date(2026, 9, 18)

    def lines(self, **over):
        base = dict(cluster="Тверь", plan=40, fact=35, draft_qty=35)
        base.update(over)
        return [SupplyLine(**base)]

    def test_valid(self):
        self.assertEqual(validate_before_create("A", self.lines(), self.ok_slot, self.now, False), [])
        self.assertEqual(validate_before_create("A", self.lines(plan=40, fact=40, draft_qty=40), self.ok_slot, self.now, False), [])
        self.assertEqual(validate_before_create("A", self.lines(plan=8, fact=8, draft_qty=8), self.ok_slot, self.now, False), [])
        self.assertEqual(validate_before_create("A", self.lines(fact=0, draft_qty=None), self.ok_slot, self.now, False), [])

    def test_fact_above_plan(self):
        errs = validate_before_create("A", self.lines(fact=45, draft_qty=45), self.ok_slot, self.now, False)
        self.assertTrue(any("больше плана" in e for e in errs), errs)

    def test_not_multiple_of_five(self):
        errs = validate_before_create("A", self.lines(fact=33, draft_qty=33), self.ok_slot, self.now, False)
        self.assertTrue(any("не кратен" in e for e in errs), errs)

    def test_below_ten(self):
        errs = validate_before_create("A", self.lines(plan=40, fact=5, draft_qty=5), self.ok_slot, self.now, False)
        self.assertTrue(any("меньше 10" in e for e in errs), errs)

    def test_draft_mismatch(self):
        errs = validate_before_create("A", self.lines(draft_qty=30), self.ok_slot, self.now, False)
        self.assertTrue(any("в черновике" in e for e in errs), errs)
        errs = validate_before_create("A", self.lines(draft_qty=None), self.ok_slot, self.now, False)
        self.assertTrue(any("в черновике" in e for e in errs), errs)

    def test_slot_too_early(self):
        errs = validate_before_create("A", self.lines(), date(2026, 9, 17), self.now, False)
        self.assertTrue(any("раньше" in e for e in errs), errs)
        errs = validate_before_create("A", self.lines(), None, self.now, False)
        self.assertTrue(any("не выбран" in e for e in errs), errs)

    def test_already_created(self):
        errs = validate_before_create("A", self.lines(), self.ok_slot, self.now, True)
        self.assertTrue(any("уже создана" in e for e in errs), errs)


class Sheet(unittest.TestCase):
    """Чтение основной раскладки (заголовок в строке 3, A/B/T/U/V) и запись результата в копию."""

    @staticmethod
    def make_book() -> bytes:
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Поставка 11.09"
        ws["A1"] = "Заявки на поставки"
        ws["A3"], ws["B3"], ws["T3"], ws["U3"] = "Артикул", "Кластер", "Заявку на поставку", "факт"
        ws["H3"], ws["S3"] = "В заявках на поставку", "Добавить за вычетом того, что уже в заявке"  # ловушки
        ws["H4"] = 99
        data = [
            (4, "AT-P-P3", "Тверь", 10), (5, "AT-P-P3", "Казань", 40), (6, "AT-P-P3", "Пермь", 20),
            (7, "AT-L-A4", "Уфа", 8), (8, "AT-L-A4", "Омск", None), (9, "AT-L-A4", "Тверь", 15),
        ]
        for r, a, b, t in data:
            ws.cell(r, 1, a)
            ws.cell(r, 2, b)
            if t is not None:
                ws.cell(r, 20, float(t))
        ws["V5"] = 70            # итог по артикулу — не трогать
        ws["C4"] = "служебное"
        ws.row_dimensions[6].hidden = True
        for col in "CDEFGHIJLNOPQR":
            ws.column_dimensions[col].hidden = True
        ws.freeze_panes = "A4"
        buffer = io.BytesIO()
        wb.save(buffer)
        return buffer.getvalue()

    def test_read_and_write(self):
        source = self.make_book()
        before = bytes(source)
        wb, ws, layout, rows = load_sheet(source, "поставка 11.09")
        self.assertEqual((layout.header_row, layout.col_article, layout.col_cluster), (3, 1, 2))
        self.assertEqual((layout.col_plan, layout.col_fact, layout.col_note), (20, 21, 23))
        self.assertEqual([(r.row, r.article, r.cluster, r.plan, r.hidden) for r in rows], [
            (4, "AT-P-P3", "Тверь", 10, False), (5, "AT-P-P3", "Казань", 40, False),
            (6, "AT-P-P3", "Пермь", 20, True), (7, "AT-L-A4", "Уфа", 8, False),
            (9, "AT-L-A4", "Тверь", 15, False)])

        out = write_results(wb, ws, layout, {4: 10, 5: 0, 7: 8}, {4: "черновик 1", 5: "исключён", 7: "x"})
        self.assertEqual(source, before, "исходные байты не должны меняться")

        import openpyxl
        ws2 = openpyxl.load_workbook(io.BytesIO(out))["Поставка 11.09"]
        self.assertEqual((ws2["U4"].value, ws2["U5"].value, ws2["U7"].value, ws2["U9"].value), (10, 0, 8, None))
        self.assertEqual(ws2["U5"].fill.fgColor.rgb, cfgmod.RED_FILL)
        self.assertNotEqual(ws2["U4"].fill.fgColor.rgb, cfgmod.RED_FILL)
        self.assertEqual(ws2["W3"].value, cfgmod.NOTE_HEADER)
        self.assertEqual((ws2["W4"].value, ws2["W5"].value), ("черновик 1", "исключён"))
        self.assertEqual(ws2["V5"].value, 70)
        self.assertTrue(ws2.row_dimensions[6].hidden)
        self.assertTrue(ws2.column_dimensions["C"].hidden)
        self.assertEqual(ws2.freeze_panes, "A4")

    def test_small_layout(self):
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "10.09 максы"
        ws.append(["Артикул", "Кластер", "Заявку на поставку"])
        ws.append(["AT-L-AM-S", "Москва, МО и Дальние регионы", 8])
        ws.append(["AT-L-AM-S", "Ростов", "12"])
        buffer = io.BytesIO()
        wb.save(buffer)
        _, _, layout, rows = load_sheet(buffer.getvalue(), "10.09 максы")
        self.assertEqual((layout.header_row, layout.col_plan, layout.col_fact, layout.col_note), (1, 3, 4, 5))
        self.assertEqual([(r.cluster, r.plan) for r in rows], [("Москва, МО и Дальние регионы", 8), ("Ростов", 12)])


def _bare_api() -> OzonFboApi:
    """Клиент без HTTP-сессии: тесты подменяют post/timeslots."""
    api = OzonFboApi.__new__(OzonFboApi)
    api.log = RunLog(None)
    api._last = 0.0
    api._calls = {}
    api.calls = []
    return api


class DraftRetry(unittest.TestCase):
    """Временный отказ NOT_AVAILABLE_CLUSTERS: одна повторная попытка через паузу."""

    def make_api(self, responses):
        api = _bare_api()
        api.post = lambda path, body, name: (api.calls.append((path, name)), responses.pop(0))[1]
        return api

    def test_retry_once_then_success(self):
        refusal = {"draft_id": 0, "errors": [{"error_message": "CAN_NOT_CREATE_DRAFT", "error_reasons": ["NOT_AVAILABLE_CLUSTERS"]}]}
        api = self.make_api([refusal, {"draft_id": 5, "errors": []}])
        slept = []
        with mock.patch("time.sleep", slept.append):
            draft_id, data = api.create_draft(1, {4070: 10}, {"warehouse_id": 1, "warehouse_type": "DELIVERY_POINT"}, "x")
        self.assertEqual((draft_id, data["draft_id"]), (5, 5))
        self.assertEqual([n for _, n in api.calls], ["x", "x_retry"])
        self.assertEqual(slept, [FboSettings().unavailable_retry_sec])
        self.assertTrue(any("повторю через" in line for line in api.log.report_lines))

    def test_no_retry_for_other_errors_or_success(self):
        api = self.make_api([{"draft_id": 7, "errors": []}])
        self.assertEqual(api.create_draft(1, {4070: 10}, {}, "ok")[0], 7)
        self.assertEqual(len(api.calls), 1)
        other = {"draft_id": 0, "errors": [{"error_message": "CAN_NOT_CREATE_DRAFT", "error_reasons": ["SOMETHING_ELSE"]}]}
        api = self.make_api([other])
        self.assertEqual(api.create_draft(1, {4070: 10}, {}, "no")[0], 0)
        self.assertEqual(len(api.calls), 1)


class RateLimit(unittest.TestCase):
    """Отдельный лимит на создание черновиков: не больше 2 успешных вызовов в окне."""

    def test_wait_for_limit(self):
        api = _bare_api()
        path = "/v1/draft/multi-cluster/create"
        slept = []
        with mock.patch("time.sleep", lambda s: slept.append(s)):
            api._wait_for_limit(path)                  # ничего не было — без ожидания
            self.assertEqual(slept, [])
            now = time.monotonic()
            api._calls[path] = [now - 50, now - 10]     # два успешных за последнюю минуту
            api._wait_for_limit(path)
            self.assertEqual(len(slept), 1)
            self.assertAlmostEqual(slept[0], 12, delta=1.5)   # ждём, пока старшему не исполнится 62 с
            self.assertTrue(any("лимит Ozon" in line for line in api.log.report_lines))
            api._calls[path] = [now - 70, now - 10]     # старший уже вне окна
            slept.clear()
            api._wait_for_limit(path)
            self.assertEqual(slept, [])
            api._wait_for_limit("/v1/supply-order/bundle")   # у других методов лимита нет
            self.assertEqual(slept, [])


class HttpClient(unittest.TestCase):
    """post(): заголовки без утечки в журнал, разбор ошибок Ozon, повтор при лимите — через httpx.MockTransport."""

    def test_post_errors_and_retry(self):
        import httpx
        seen: list[httpx.Request] = []
        responses = [
            httpx.Response(429, json={"code": 8, "message": "You have reached request rate limit per second"}),
            httpx.Response(200, json={"result": [1, 2]}),
            httpx.Response(200, json={"code": 5, "message": "Invalid Api-Key, please check the key and try again"}),
            httpx.Response(400, json={"code": 3, "message": "Request validation error"}),
            httpx.Response(200, text="not json"),
        ]

        def handler(request: httpx.Request) -> httpx.Response:
            seen.append(request)
            return responses.pop(0)

        log = RunLog(None)
        api = OzonFboApi("client", "secret-key", log, transport=httpx.MockTransport(handler))
        with mock.patch("time.sleep", lambda s: None):
            self.assertEqual(api.post("/v2/cluster/list", {"a": 1}, "x"), {"result": [1, 2]})
            self.assertEqual(seen[0].headers["Client-Id"], "client")
            self.assertEqual(seen[0].headers["Api-Key"], "secret-key")
            self.assertTrue(any("лимит запросов Ozon" in line for line in log.report_lines))
            with self.assertRaises(api_client.AuthError):
                api.post("/v2/cluster/list", {}, "auth")
            with self.assertRaises(ApiError) as ctx:
                api.post("/v2/cluster/list", {}, "bad")
            self.assertEqual(ctx.exception.status, 400)
            self.assertEqual(api.post("/v2/cluster/list", {}, "raw"), {"raw": "not json"})
        self.assertFalse(any("secret-key" in line for line in log.report_lines), "ключ не попадает в журнал")
        api.close()

    def test_network_error_is_wrapped(self):
        import httpx

        def handler(request: httpx.Request) -> httpx.Response:
            raise httpx.ConnectError("boom", request=request)

        api = OzonFboApi("c", "k", RunLog(None), transport=httpx.MockTransport(handler))
        with mock.patch("time.sleep", lambda s: None), self.assertRaises(api_client.NetworkError):
            api.post("/v2/cluster/list", {}, "x")
        api.close()

    def test_credentials_prefer_fbo_specific(self):
        with mock.patch.dict("os.environ", {"OZON_CLIENT_ID": "a", "OZON_API_KEY": "b"}):
            self.assertEqual(api_client.get_credentials(), ("a", "b"))
            self.assertTrue(api_client.credentials_configured())
        with mock.patch.dict("os.environ", {"OZON_CLIENT_ID": "a", "OZON_API_KEY": "b",
                                            "OZON_FBO_CLIENT_ID": "fbo", "OZON_FBO_API_KEY": "fbo-key"}):
            self.assertEqual(api_client.get_credentials(), ("fbo", "fbo-key"))
        with mock.patch.dict("os.environ", {"OZON_CLIENT_ID": "", "OZON_API_KEY": ""}):
            self.assertFalse(api_client.credentials_configured())


class FakeApi:
    """Подменённый Ozon для расчёта артикула: подбор по кластерам, общий черновик, слот."""

    def __init__(self, capacity, slots_days, combined_qty=None, days_fn=None, states_fn=None):
        self.capacity = capacity
        self.days = slots_days
        self.combined_qty = combined_qty or {}
        self.days_fn = days_fn          # (состав последнего черновика) → дни со слотами
        self.states_fn = states_fn      # (номер черновика, кластер, состав) → (состояние склада, кол-во в составе или None)
        self.drafts = 0
        self.compositions = []          # составы всех созданных черновиков
        self.log = RunLog(None)

    def create_draft(self, sku, cluster_qty, drop_off, name):
        self.drafts += 1
        self.last = dict(cluster_qty)
        self.last_sku = sku
        self.by_draft = getattr(self, "by_draft", {})
        self.by_draft[self.drafts] = (sku, dict(cluster_qty))
        self.compositions.append(dict(cluster_qty))
        return self.drafts, {"draft_id": self.drafts, "errors": []}

    def _warehouse(self, cid):
        state = self.states_fn(self.drafts, cid, dict(self.last))[0] if self.states_fn else "FULL_AVAILABLE"
        if state == "NOT_AVAILABLE":
            return {"restricted_bundle_id": f"r{cid}",
                    "availability_status": {"state": state, "invalid_reason": "NOT_AVAILABLE_TIMESLOT_FOR_BOTH_WAREHOUSES"}}
        return {"bundle_id": f"b{cid}", "availability_status": {"state": state, "invalid_reason": "UNSPECIFIED"}}

    def wait_draft(self, draft_id, name, tries=20):
        return {"status": "SUCCESS", "errors": [], "clusters": [
            {"macrolocal_cluster_id": cid, "warehouses": [self._warehouse(cid)]} for cid in self.last]}

    def bundle_qty(self, bundle_id, sku, name):
        cid = int(bundle_id[1:])
        asked = self.last[cid]
        if self.states_fn:
            q = self.states_fn(self.drafts, cid, dict(self.last))[1]
            if q is not None:
                return q
        if cid in self.combined_qty and len(self.last) > 1:
            return self.combined_qty[cid]
        return min(asked, self.capacity.get(cid, asked))

    def timeslots(self, draft_id, cluster_ids, today, name, date_from=None):
        days = self.days_fn(self.last) if self.days_fn else self.days
        return {"error_reason": "UNSPECIFIED", "result": {"drop_off_warehouse_timeslots": {"days": days}}}

    def timeslot_days(self, draft_id, cluster_ids, today, name):
        return (self.days_fn(self.last) if self.days_fn else self.days), None


NOW = datetime(2026, 9, 11, 12, 0, tzinfo=MSK)
DAYS = [{"date_in_timezone": "2026-09-25", "timeslots": [
    {"from_in_timezone": "2026-09-25T18:00:00", "to_in_timezone": "2026-09-25T19:00:00"},
    {"from_in_timezone": "2026-09-25T20:00:00", "to_in_timezone": "2026-09-25T21:00:00"}]}]
EARLY = [{"date_in_timezone": "2026-09-15", "timeslots": [{"from_in_timezone": "2026-09-15T10:00:00", "to_in_timezone": "2026-09-15T11:00:00"}]}]


class ArticleFlow(unittest.TestCase):
    """process_article с подменённым API: подбор по кластерам, общий черновик, слот, примечания."""

    def lines(self):
        return [ClusterLine(4, "Тверь", 4072, 40), ClusterLine(5, "Казань", 4041, 20),
                ClusterLine(6, "Луна", None, 10)]

    def lines3(self):
        return [ClusterLine(4, "Тверь", 4072, 40), ClusterLine(5, "Казань", 4041, 20), ClusterLine(6, "Калининград", 4004, 10)]

    def test_trial_flow(self):
        api = FakeApi({4072: 23, 4041: 9}, DAYS)
        out = []
        res = process_article(api, "AT-P-P3", self.lines(), 1891251877, {}, NOW, out.append)
        self.assertEqual([ln.fact for ln in res.lines], [20, 0, None])
        self.assertEqual(res.lines[1].note[:9], "исключён:")
        self.assertTrue(any("Луна" in r and "не найден" in r for r in res.manual), res.manual)
        self.assertEqual(api.compositions, [{4072: 40, 4041: 20}, {4072: 20}, {4072: 20}])
        self.assertEqual(res.slot["from"], "2026-09-25T20:00:00")
        notes = outcome_rows(res, "trial")[1]
        self.assertTrue(notes[4].startswith("РУЧНАЯ ПРОВЕРКА"))
        self.assertTrue(notes[5].startswith("исключён"))
        self.assertIn("не найден", notes[6])

    def test_all_accepted_needs_one_draft(self):
        api = FakeApi({}, DAYS)
        res = process_article(api, "AT-P-P3", self.lines()[:2], 1, {}, NOW, lambda s: None)
        self.assertEqual([ln.fact for ln in res.lines], [40, 20])
        self.assertEqual([ln.draft_qty for ln in res.lines], [40, 20])
        self.assertEqual(api.compositions, [{4072: 40, 4041: 20}], "если Ozon принял весь план — один черновик")
        self.assertEqual((res.draft_id, res.manual), (1, []))

    def test_plan_ok_but_no_slot_falls_back_to_singles(self):
        api = FakeApi({}, DAYS, days_fn=lambda last: EARLY if last.get(4041, 0) >= 20 else DAYS)
        res = process_article(api, "AT-P-P3", self.lines()[:2], 1, {}, NOW, lambda s: None)
        self.assertEqual([ln.fact for ln in res.lines], [40, 15])
        self.assertEqual(res.manual, [])
        self.assertEqual(api.compositions[0], {4072: 40, 4041: 20})
        self.assertEqual(api.compositions[-1], {4072: 40, 4041: 15})
        self.assertEqual(res.slot["from"], "2026-09-25T20:00:00")

    def test_final_slot_lost_rechecks_provisional(self):
        api = FakeApi({4072: 23}, DAYS, days_fn=lambda last: EARLY if last.get(4041, 0) >= 20 else DAYS)
        res = process_article(api, "AT-P-P3", self.lines()[:2], 1, {}, NOW, lambda s: None)
        self.assertEqual([ln.fact for ln in res.lines], [20, 15])
        self.assertEqual(res.manual, [])
        self.assertEqual(api.compositions[-1], {4072: 20, 4041: 15})

    def test_clean_flow_notes(self):
        api = FakeApi({}, DAYS)
        res = process_article(api, "AT-P-P3", self.lines()[:2], 1, {}, NOW, lambda s: None)
        self.assertEqual([ln.fact for ln in res.lines], [40, 20])
        self.assertEqual(res.manual, [])
        self.assertEqual(res.draft_id, 1)
        notes = outcome_rows(res, "trial")[1]
        self.assertIn("заявка не создана", notes[4].lower())
        self.assertIn("2026-09-25 20:00", notes[4])
        errs = validate_before_create("AT-P-P3", [SupplyLine(l.cluster, l.plan, l.fact, l.draft_qty) for l in res.lines],
                                      date.fromisoformat(res.slot["date"]), NOW, False)
        self.assertEqual(errs, [])

    def test_final_unavailable_retry_then_exclude(self):
        api = FakeApi({4041: 15}, DAYS,
                      states_fn=lambda n, cid, comp: ("NOT_AVAILABLE" if cid == 4004 and n >= 2 else "FULL_AVAILABLE", None))
        out = []
        res = process_article(api, "AT-P-P3", self.lines3(), 1, {}, NOW, out.append)
        self.assertEqual([ln.fact for ln in res.lines], [40, 15, 0])
        self.assertEqual(api.compositions, [{4072: 40, 4041: 20, 4004: 10}, {4041: 15},
                                            {4072: 40, 4041: 15, 4004: 10}, {4072: 40, 4041: 15, 4004: 10},
                                            {4072: 40, 4041: 15}])
        self.assertTrue(res.lines[2].note.startswith("исключён в итоговом черновике"), res.lines[2].note)
        self.assertIn("нет таймслотов для складов кластера", res.lines[2].note)
        self.assertEqual(res.manual, [])
        self.assertIsNotNone(res.slot)
        self.assertEqual([ln.draft_qty for ln in res.lines], [40, 15, None])
        self.assertTrue(any("повторяю тот же состав" in s for s in out), out)
        self.assertTrue(any("исключаю из поставки" in s for s in out), out)
        errs = validate_before_create("AT-P-P3", [SupplyLine(l.cluster, l.plan, l.fact or 0, l.draft_qty) for l in res.lines],
                                      date.fromisoformat(res.slot["date"]), NOW, False)
        self.assertEqual(errs, [])

    def test_final_unavailable_transient(self):
        api = FakeApi({4041: 15}, DAYS,
                      states_fn=lambda n, cid, comp: ("NOT_AVAILABLE" if cid == 4004 and n == 3 else "FULL_AVAILABLE", None))
        res = process_article(api, "AT-P-P3", self.lines3(), 1, {}, NOW, lambda s: None)
        self.assertEqual([ln.fact for ln in res.lines], [40, 15, 10])
        self.assertEqual(res.manual, [])
        self.assertEqual(api.compositions[-2:], [{4072: 40, 4041: 15, 4004: 10}, {4072: 40, 4041: 15, 4004: 10}])
        self.assertIsNotNone(res.slot)

    def test_final_trimmed_repicks_quantity(self):
        def states(n, cid, comp):
            if cid == 4072 and len(comp) > 1 and n >= 3 and comp[cid] > 20:
                return "FULL_AVAILABLE", 23
            return "FULL_AVAILABLE", None
        api = FakeApi({4041: 15}, DAYS, states_fn=states)
        res = process_article(api, "AT-P-P3", self.lines3(), 1, {}, NOW, lambda s: None)
        self.assertEqual([ln.fact for ln in res.lines], [20, 15, 10])
        self.assertEqual(res.manual, [])
        self.assertEqual(api.compositions[-1], {4072: 20, 4041: 15, 4004: 10})
        self.assertIsNotNone(res.slot)

    def test_final_rebuild_limit_gives_manual_or_empty(self):
        api = FakeApi({4041: 15}, DAYS,
                      states_fn=lambda n, cid, comp: ("NOT_AVAILABLE" if n >= 2 and cid == min(comp) else "FULL_AVAILABLE", None))
        res = process_article(api, "AT-P-P3", self.lines3(), 1, {}, NOW, lambda s: None)
        self.assertLessEqual(sum(1 for c in api.compositions if len(c) > 1), 1 + cfgmod.FINAL_REBUILD_LIMIT)
        self.assertTrue(res.manual or res.note, (res.manual, res.note))
        self.assertIsNone(res.slot)

    def test_combined_mismatch_is_manual(self):
        api = FakeApi({}, DAYS, combined_qty={4072: 33})
        res = process_article(api, "AT-P-P3", self.lines()[:2], 1, {}, NOW, lambda s: None)
        self.assertEqual([ln.fact for ln in res.lines], [30, 20])
        self.assertTrue(any("33 шт. вместо 30" in r for r in res.manual), res.manual)
        self.assertIsNone(res.slot)

    def test_no_slot_is_manual(self):
        api = FakeApi({}, EARLY)
        res = process_article(api, "AT-P-P3", self.lines()[:2], 1, {}, NOW, lambda s: None,
                              check_fn=lambda *a, **k: CheckResult(True))
        self.assertTrue(any("нет слота" in r for r in res.manual), res.manual)

    def test_all_excluded(self):
        api = FakeApi({4072: 3, 4041: 3}, DAYS)
        res = process_article(api, "AT-P-P3", self.lines()[:2], 1, {}, NOW, lambda s: None)
        self.assertEqual([ln.fact for ln in res.lines], [0, 0])
        self.assertIn("не создаётся", res.note)
        self.assertEqual(res.draft_id, 0)


class FakeLiveApi(FakeApi):
    """Одна заявка на артикул, внутри supplies[] по поставке на кластер (как у Ozon)."""

    offer_by_sku = {1: "AT-P-P3", 1891251877: "AT-P-P3", 1716053944: "AT-L-A3"}

    def __init__(self, *a, **k):
        super().__init__(*a, **k)
        self.supply_calls = []
        self.status_response = None
        self.supply_clusters = []
        self.cargo_calls = []
        self.cargo_status = None
        self.today_orders = {}       # order_id → заявка как в /v3/supply-order/get (для проверки дублей)
        self.today_bundles = {}      # bundle_id → состав
        self.list_calls = []

    def supply_order_ids(self, states, name):
        self.list_calls.append(list(states))
        return [oid for oid, o in self.today_orders.items() if o.get("state") in states]

    def bundle_items(self, bundle_id, name):
        if bundle_id in self.today_bundles:
            return [dict(i) for i in self.today_bundles[bundle_id]]
        if str(bundle_id).startswith("b") and str(bundle_id)[1:].isdigit():
            cid = int(bundle_id[1:])
            sku, comp = getattr(self, "by_draft", {}).get(getattr(self, "supply_draft", None), (getattr(self, "last_sku", 1), self.last))
            return [{"offer_id": self.offer_by_sku.get(sku, "AT-P-P3"), "sku": sku, "quantity": comp.get(cid, 0)}]
        return []

    def create_supply(self, draft_id, cluster_ids, slot, name):
        self.supply_calls.append((draft_id, list(cluster_ids), dict(slot)))
        self.supply_clusters = list(cluster_ids)
        self.supply_draft = draft_id
        return {"error_reasons": [], "draft_id": draft_id}

    def supply_create_status(self, draft_id, operation_id, name):
        return self.status_response or {"order_id": 900 + draft_id, "status": "SUCCESS", "error_reasons": []}

    def supply_orders(self, order_ids, name):
        if order_ids and all(int(i) in self.today_orders for i in order_ids):
            return {"orders": [self.today_orders[int(i)] for i in order_ids]}
        return {"orders": [{"order_id": i, "order_number": f"N{i}", "state": "DATA_FILLING",
                            "data_filling_deadline": "2026-09-13T09:00:00Z",
                            "supplies": [{"supply_id": 5000 + cid, "state": "DATA_FILLING",
                                          "macrolocal_cluster_id": str(cid), "bundle_id": f"b{cid}"}
                                         for cid in self.supply_clusters]} for i in order_ids]}

    def cargo_rules(self, supply_ids, name):
        return {"rules": [{"supply_id": s, "max_cargo_count": 100} for s in supply_ids]}

    def create_cargoes(self, supply_id, cargoes, name):
        self.cargo_calls.append((supply_id, cargoes))
        return {"operation_id": f"op-{supply_id}"}

    def cargoes_create_info(self, operation_id, name):
        if self.cargo_status:
            return self.cargo_status
        supply_id = int(operation_id.split("-")[1])
        cargoes = next(c for s, c in self.cargo_calls if s == supply_id)
        return {"status": "SUCCESS", "errors": [],
                "cargoes": [{"key": c["key"], "cargo_id": 7000 + n} for n, c in enumerate(cargoes, 1)]}


class LiveFlow(unittest.TestCase):
    """Боевой режим: предохранители → подтверждение → создание заявки (Ozon подменён, сети нет)."""

    def ready(self, api, lines=None):
        lines = lines or [ClusterLine(4, "Тверь", 4072, 40), ClusterLine(5, "Казань", 4041, 20)]
        return process_article(api, "AT-P-P3", lines, 1, {}, NOW, lambda s: None, mode="live")

    def run_live(self, api, res, confirm=lambda text: True, already=False, with_cargo=True):
        """Заявка по одному артикулу через общий конвейер create_orders_batch с вопросом по артикулу."""
        store = InMemoryStateStore()
        if already:     # в памяти уже есть действующая заявка по этому артикулу
            save_state(store, "тест", ArticleResult(res.article, res.sku, [], orders=["OLD-1"], created=True,
                                                    order_details=[{"order_id": 901, "number": "OLD-1"}]), "live")
        out = []
        with mock.patch("time.sleep", lambda s: None):
            create_orders_batch(api, store, [res], "тест", NOW, out.append, with_cargo=with_cargo, confirm_each=confirm)
        return out

    def test_creates_order_without_cargo(self):
        api = FakeLiveApi({}, DAYS)
        res = self.ready(api)
        shown = []
        self.run_live(api, res, confirm=lambda text: shown.append(text) or True, with_cargo=False)
        self.assertTrue(res.created)
        self.assertEqual(res.orders, ["N901"])
        self.assertEqual(api.supply_calls, [(1, [4072, 4041], res.slot)])
        self.assertEqual(res.order_details[0]["supply_ids"], [9072, 9041])
        self.assertEqual(api.cargo_calls, [])
        self.assertIn("Тверь: 40 шт.", shown[0])
        self.assertIn("не передаются", shown[0])
        self.assertIn("2026-09-25 20:00", shown[0])
        note = outcome_rows(res, "live")[1][4]
        self.assertIn("заявка №N901", note)
        self.assertIn("грузоместа не переданы", note)
        self.assertIn("13.09.2026 12:00 МСК", note)

    def test_creates_order_with_cargoes(self):
        api = FakeLiveApi({}, DAYS)
        res = self.ready(api, [ClusterLine(4, "Тверь", 4072, 110), ClusterLine(5, "Казань", 4041, 20)])
        shown = []
        self.run_live(api, res, confirm=lambda text: shown.append(text) or True)
        self.assertTrue(res.created)
        self.assertIn("Тверь: 3 короба — 50+50+10 шт.", shown[0])
        self.assertIn("Казань: 1 короб — 20 шт.", shown[0])
        self.assertEqual(layout_text("AT-P-P3", 250), "5 коробов — 50+50+50+50+50 шт.")
        self.assertEqual(layout_text("AT-L-AM-S", 12), "3 короба — 5+5+2 шт.")
        self.assertEqual([(s, [cargo_qty(c) for c in cs]) for s, cs in api.cargo_calls],
                         [(9072, [50, 50, 10]), (9041, [20])])
        for _, cargoes in api.cargo_calls:
            for c in cargoes:
                self.assertEqual(c["value"]["type"], "BOX")
                self.assertEqual(c["value"]["items"][0]["offer_id"], "AT-P-P3")
                self.assertEqual(c["value"]["items"][0]["quant"], 1)
            self.assertEqual(len({c["key"] for c in cargoes}), len(cargoes), "ключи уникальны")
        self.assertEqual(res.cargoes["9072"]["cargo_ids"], [7001, 7002, 7003])
        self.assertEqual(res.cargoes["9041"]["layout"], [20])
        self.assertEqual(res.manual, [])
        note = outcome_rows(res, "live")[1][4]
        self.assertIn("грузоместа переданы: 4", note)

    def test_cargo_failure_keeps_order(self):
        api = FakeLiveApi({}, DAYS)
        api.cargo_status = {"status": "FAILED", "errors": [{"message": "bad box"}]}
        res = self.ready(api)
        out = self.run_live(api, res)
        self.assertTrue(res.created)
        self.assertEqual(res.orders, ["N901"])
        self.assertEqual(res.cargoes, {})
        self.assertTrue(any("грузоместа не переданы" in r and "bad box" in r for r in res.manual), res.manual)
        self.assertTrue(any("Дозаполнить" in line for line in out), out)
        self.assertIn("грузоместа не переданы", outcome_rows(res, "live")[1][4])

    def test_build_and_validate_order_cargoes(self):
        cargoes = build_cargoes("AT-L-AM-S", 1, 12)
        self.assertEqual([cargo_qty(c) for c in cargoes], [5, 5, 2])
        supplies = [{"supply_id": 1, "cluster_id": None, "storage": "ТВЕРЬ_РФЦ", "state": "DATA_FILLING"},
                    {"supply_id": 2, "cluster_id": None, "storage": "КАЗАНЬ_РФЦ_НОВЫЙ", "state": "DATA_FILLING"}]
        items = {1: [{"offer_id": "AT-P-P3", "sku": 1, "quantity": 110}], 2: [{"offer_id": "AT-P-P3", "sku": 1, "quantity": 20}]}
        ok = {1: build_item_cargoes(items[1]), 2: build_item_cargoes(items[2])}
        self.assertEqual(validate_order_cargoes("AT-P-P3", {4072: 110, 4041: 20}, supplies, items, ok), [])
        errs = validate_order_cargoes("AT-P-P3", {4072: 110, 4041: 20, 4070: 10}, supplies, items, ok)
        self.assertTrue(any("не совпадают с фактами [10, 20, 110]" in e for e in errs), errs)
        errs = validate_order_cargoes("AT-P-P3", {4072: 100, 4041: 20}, supplies, items, ok)
        self.assertTrue(any("[20, 110] не совпадают с фактами [20, 100]" in e for e in errs), errs)
        big = {1: [{"key": "k", "value": {"type": "BOX", "items": [{"offer_id": "AT-P-P3", "quantity": 110, "quant": 1}]}}], 2: ok[2]}
        errs = validate_order_cargoes("AT-P-P3", {4072: 110, 4041: 20}, supplies, items, big)
        self.assertTrue(any("вместимости 50" in e for e in errs), errs)
        foreign = {1: [{"offer_id": "AT-X", "sku": 9, "quantity": 110}], 2: items[2]}
        errs = validate_order_cargoes("AT-P-P3", {4072: 110, 4041: 20}, supplies, foreign, ok)
        self.assertTrue(any("чужой артикул" in e for e in errs), errs)
        cancelled = [dict(supplies[0], state="CANCELLED"), supplies[1]]
        self.assertEqual(validate_order_cargoes("AT-P-P3", {4072: 110, 4041: 20}, cancelled, items, {2: ok[2]}), [])
        errs = validate_order_cargoes("AT-P-P3", {4072: 110, 4041: 20}, [dict(supplies[0], state="READY_TO_SUPPLY"), supplies[1]],
                                      items, {2: ok[2]})
        self.assertEqual(errs, [], "уже заполненная поставка тоже пропускается")

    def test_cargoes_when_ozon_assigns_warehouses_immediately(self):
        api = FakeLiveApi({}, DAYS)
        real = api.supply_orders

        def no_cluster(order_ids, name):
            data = real(order_ids, name)
            for o in data["orders"]:
                for s in o["supplies"]:
                    s["storage_warehouse"] = {"warehouse_id": 10 + int(s["macrolocal_cluster_id"]), "name": "СКЛАД_" + s["macrolocal_cluster_id"]}
                    s["macrolocal_cluster_id"] = None
            return data
        api.supply_orders = no_cluster
        res = self.ready(api, [ClusterLine(4, "Тверь", 4072, 110), ClusterLine(5, "Казань", 4041, 20)])
        self.run_live(api, res)
        self.assertTrue(res.created)
        self.assertEqual(res.manual, [])
        self.assertEqual(sorted((s, [cargo_qty(c) for c in cs]) for s, cs in api.cargo_calls), [(9041, [20]), (9072, [50, 50, 10])])
        self.assertEqual(res.cargoes["9072"]["storage"], "СКЛАД_4072")
        self.assertEqual(res.warehouses, {"поставка 9072": "СКЛАД_4072", "поставка 9041": "СКЛАД_4041"})
        res.wh_map = {4082: 4072, 4051: 4041}
        self.assertEqual(order_warehouses(res), {"Тверь": "СКЛАД_4072", "Казань": "СКЛАД_4041"})

    def test_declined_by_user(self):
        api = FakeLiveApi({}, DAYS)
        res = self.ready(api)
        self.run_live(api, res, confirm=lambda text: False)
        self.assertEqual(api.supply_calls, [])
        self.assertFalse(res.created)
        self.assertIn("отменено", res.note)

    def test_already_created_blocks(self):
        api = FakeLiveApi({}, DAYS)
        res = self.ready(api)
        self.run_live(api, res, already=True)
        self.assertEqual(api.supply_calls, [])
        self.assertIn("уже создана: OLD-1", res.note)
        self.assertFalse(res.created)

    def test_manual_article_not_created(self):
        api = FakeLiveApi({}, DAYS)
        res = self.ready(api, [ClusterLine(4, "Тверь", 4072, 40), ClusterLine(6, "Луна", None, 10)])
        self.assertTrue(res.manual)
        self.run_live(api, res)
        self.assertEqual(api.supply_calls, [])

    def test_slot_too_early_blocks(self):
        api = FakeLiveApi({}, DAYS)
        res = self.ready(api)
        res.slot = {"date": "2026-09-15", "from": "2026-09-15T10:00:00", "to": "2026-09-15T11:00:00"}
        self.run_live(api, res)
        self.assertEqual(api.supply_calls, [])
        self.assertTrue(any("раньше" in r for r in res.manual), res.manual)

    def test_status_without_order_is_manual(self):
        api = FakeLiveApi({}, DAYS)
        api.status_response = {"status": "ERROR", "error_messages": ["boom"]}
        res = self.ready(api)
        self.run_live(api, res)
        self.assertEqual(len(api.supply_calls), 1)
        self.assertFalse(res.created)
        self.assertTrue(any("создание заявки" in r and "boom" in r for r in res.manual), res.manual)

    def test_cluster_unavailable_at_draft_creation(self):
        refusal = {"draft_id": 0, "errors": [{"error_message": "CAN_NOT_CREATE_DRAFT", "macrolocal_cluster_ids": [],
                                              "skus": [], "error_reasons": ["NOT_AVAILABLE_CLUSTERS"],
                                              "message": "", "items_validation": []}]}

        class RefusingApi(FakeLiveApi):
            def create_draft(self, sku, cluster_qty, drop_off, name):
                self.compositions.append(dict(cluster_qty))
                return 0, refusal

        api = RefusingApi({}, DAYS)
        res = self.ready(api, [ClusterLine(4, "Пермь", 4070, 40)])
        self.assertEqual([ln.fact for ln in res.lines], [0])
        self.assertEqual(api.compositions[:2], [{4070: 40}, {4070: 40}])
        self.assertEqual(api.compositions[-1], {4070: 10})
        self.assertIn("недоступен", res.lines[0].note)
        self.assertIn("NOT_AVAILABLE_CLUSTERS", res.lines[0].note)
        self.assertIn("не создаётся", res.note)
        self.assertEqual(res.manual, [])
        self.run_live(api, res)
        self.assertEqual(api.supply_calls, [])
        self.assertEqual(res.manual, [])
        self.assertTrue(outcome_rows(res, "live")[1][4].startswith("исключён"))

    def test_parse_helpers(self):
        self.assertEqual(parse_order_ids({"status": "SUCCESS", "result": {"order_ids": [2, 1]}}), [1, 2])
        self.assertEqual(parse_order_ids({"status": "SUCCESS", "order_id": 7}), [7])
        self.assertEqual(parse_order_ids({"status": "SUCCESS", "orders": [{"order_id": "5"}]}), [5])
        self.assertEqual(parse_order_ids({"status": "IN_PROGRESS"}), [])
        real = {"orders": [{"order_id": 128199092, "order_number": "2000066365224", "state": "DATA_FILLING",
                            "data_filling_deadline": "2026-10-09T15:00:00Z",
                            "order_tags": {"is_super_fbo": True, "original_supply_id": 0},
                            "timeslot": {"timeslot": {"from": "2026-10-09T16:00:00Z", "to": "2026-10-09T17:00:00Z"}},
                            "supplies": [{"state": "DATA_FILLING", "supply_id": 2000066365224, "storage_warehouse": None,
                                          "is_crossdock": True, "bundle_id": "01a0902c", "macrolocal_cluster_id": "4070"}]}]}
        orders = parse_orders(real)
        self.assertEqual(len(orders), 1)
        o = orders[0]
        self.assertEqual((o["order_id"], o["number"], o["state"]), (128199092, "2000066365224", "DATA_FILLING"))
        self.assertEqual(o["deadline_utc"], "2026-10-09T15:00:00Z")
        self.assertEqual(o["supply_ids"], [2000066365224])
        self.assertEqual(o["supplies"], [{"supply_id": 2000066365224, "cluster_id": 4070, "state": "DATA_FILLING", "bundle_id": "01a0902c",
                                          "storage": None, "storage_id": None}])
        self.assertEqual(utc_to_msk_text(o["deadline_utc"]), "09.10.2026 18:00 МСК")
        self.assertEqual(utc_to_msk_text("2026-09-13T09:00:00Z"), "13.09.2026 12:00 МСК")
        self.assertEqual(utc_to_msk_text(None), "?")


class TimeslotPaging(unittest.TestCase):
    """Слоты запрашиваются окнами по 28 дней вперёд, пока Ozon их возвращает; берём самый поздний."""

    def make_api(self, windows, fail_from=None):
        api = _bare_api()

        def timeslots(draft_id, cluster_ids, today, name, date_from=None):
            start = date_from or today
            api.calls.append(start)
            if fail_from is not None and start >= fail_from:
                raise ApiError("/v2/draft/timeslot/info", 400, {"code": 3, "message": "too far"})
            days = windows.get(start, [])
            return {"error_reason": "UNSPECIFIED", "result": {"drop_off_warehouse_timeslots": {"days": days}}}
        api.timeslots = timeslots
        return api

    @staticmethod
    def day(d, slots=("19:00", "20:00")):
        return {"date_in_timezone": d, "timeslots": [{"from_in_timezone": f"{d}T{a}:00", "to_in_timezone": f"{d}T{b}:00"}
                                                     for a, b in [slots]]}

    def test_pages_until_empty(self):
        today = date(2026, 9, 11)
        w1 = [self.day("2026-09-11"), self.day("2026-10-09")]
        w2 = [self.day("2026-10-09"), self.day("2026-10-20"), self.day("2026-11-06", ("09:00", "10:00"))]
        w3 = [{"date_in_timezone": "2026-11-07", "timeslots": []}]
        api = self.make_api({today: w1, today + timedelta(days=28): w2, today + timedelta(days=56): w3})
        days, reason = api.timeslot_days(1, [4070], today, "s")
        self.assertIsNone(reason)
        self.assertEqual([d["date_in_timezone"] for d in days],
                         ["2026-09-11", "2026-10-09", "2026-10-20", "2026-11-06", "2026-11-07"])
        self.assertEqual(api.calls, [today, today + timedelta(days=28), today + timedelta(days=56)])
        slot = pick_latest_slot(days, date(2026, 9, 18))
        self.assertEqual((slot["date"], slot["from"]), ("2026-11-06", "2026-11-06T09:00:00"))

    def test_error_in_later_window_keeps_earlier(self):
        today = date(2026, 9, 11)
        api = self.make_api({today: [self.day("2026-10-09")]}, fail_from=today + timedelta(days=28))
        days, reason = api.timeslot_days(1, [4070], today, "s")
        self.assertEqual([d["date_in_timezone"] for d in days], ["2026-10-09"])
        self.assertIsNone(reason)

    def test_error_in_first_window_raises(self):
        today = date(2026, 9, 11)
        api = self.make_api({}, fail_from=today)
        with self.assertRaises(ApiError):
            api.timeslot_days(1, [4070], today, "s")


class Settings(unittest.TestCase):
    def test_parse_and_apply(self):
        settings, errs = parse_settings({"default_cargo_capacity": "40", "cargo_capacity": {"at-x": "7"}, "min_slot_offset_days": 0})
        self.assertEqual(errs, [])
        with use_settings(settings):
            self.assertEqual(cargo_layout("AT-P-P3", 90), [40, 40, 10])
            self.assertEqual(cargo_layout("AT-X", 15), [7, 7, 1])
            self.assertEqual(min_slot_date(datetime(2026, 9, 11, tzinfo=MSK)), date(2026, 9, 11))
        self.assertEqual(cargo_layout("AT-P-P3", 90), [50, 40], "вне контекста действуют значения по умолчанию")
        again, errs = parse_settings(settings_to_dict(settings))
        self.assertEqual((errs, again), ([], settings), "round-trip через словарь без потерь")
        with_drop, errs = parse_settings({"drop_off_name": "ТЕСТ_1"}, settings)
        self.assertEqual((errs, with_drop.drop_off_name, with_drop.default_cargo_capacity), ([], "ТЕСТ_1", 40))

    def test_capacity_text_form(self):
        settings, errs = parse_settings({"cargo_capacity": "AT-L-AM-S = 5\n\n at-y=12 "})
        self.assertEqual(errs, [])
        self.assertEqual(settings.cargo_capacity, {"AT-L-AM-S": 5, "AT-Y": 12})
        _, errs = parse_settings({"cargo_capacity": "без знака равно"})
        self.assertEqual(len(errs), 1)

    def test_bad_values_reported_and_ignored(self):
        before = FboSettings()
        settings, errs = parse_settings({"default_cargo_capacity": "много", "min_qty": 0, "cargo_capacity": {"A": -1}, "drop_off_name": " "})
        self.assertEqual(len(errs), 4, errs)
        self.assertEqual(settings, before)
        settings, errs = parse_settings({"all_clusters": "Казань\n\n Луна "})
        self.assertEqual((errs, settings.all_clusters), ([], ["Казань", "Луна"]))


class ManualRows(unittest.TestCase):
    def test_parse(self):
        text = "Артикул\tКластер\tЗаявку на поставку\nAT-L-A3\tПермь\t10\nAT-P-P3; Москва, МО и Дальние регионы; 20,0\n\nAT-P-P3;Тверь;0\n"
        rows = parse_manual_rows(text)
        self.assertEqual([(r.row, r.article, r.cluster, r.plan, r.hidden) for r in rows],
                         [(2, "AT-L-A3", "Пермь", 10, False), (3, "AT-P-P3", "Москва, МО и Дальние регионы", 20, False)])
        with self.assertRaises(ValueError):
            parse_manual_rows("AT-L-A3; Пермь")
        with self.assertRaises(ValueError):
            parse_manual_rows("AT-L-A3; Пермь; десять")
        self.assertEqual(parse_manual_rows(""), [])


if __name__ == "__main__":
    unittest.main()
